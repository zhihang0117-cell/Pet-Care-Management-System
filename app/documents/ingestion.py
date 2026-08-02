"""
Load a company policy document (.docx) from Supabase Storage for RAG
indexing. Ported from the retired ai-backend/app/rag/{ingestion,storage}.py
and app/services/ingestion.py (pre-restructure commit 2d5ecb4^).
"""

from __future__ import annotations

import tempfile
from pathlib import Path

from docx import Document as DocxDocument
from langchain_core.documents import Document

from app.db.supabase_client import get_supabase_client

PRODUCTION_DOCUMENT_TYPES = frozenset({"policies", "service_information", "business_flow_booking", "veterinary"})


def validate_document_type(document_type: str) -> str:
    normalized = str(document_type or "").strip().lower()
    if normalized not in PRODUCTION_DOCUMENT_TYPES:
        allowed = ", ".join(sorted(PRODUCTION_DOCUMENT_TYPES))
        raise ValueError(f"Unsupported document_type '{document_type}'. Allowed business categories: {allowed}")
    return normalized


def download_storage_docx(storage_bucket: str, storage_path: str) -> Path:
    bucket = str(storage_bucket or "").strip()
    path = str(storage_path or "").strip().lstrip("/")
    if not bucket:
        raise ValueError("storage_bucket is required")
    if not path:
        raise ValueError("storage_path is required")
    if not path.lower().endswith(".docx"):
        raise ValueError("storage_path must point to a .docx file")

    client = get_supabase_client()
    payload = client.storage.from_(bucket).download(path)
    if not payload:
        raise ValueError(f"Could not download object: {bucket}/{path}")

    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
    temp_file.write(payload)
    temp_file.flush()
    temp_file.close()
    return Path(temp_file.name)


def load_docx_pages(path: Path) -> list[Document]:
    """Extract text (paragraphs + table cells) from a Word document."""
    if not path.exists():
        raise FileNotFoundError(f"Document not found: {path}")
    if path.suffix.lower() != ".docx":
        raise ValueError(f"Unsupported file format: {path.name}. Upload .docx Word files only.")

    doc = DocxDocument(str(path))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    # Policies often use tables — include cell text too.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    paragraphs.append(text)

    if not paragraphs:
        raise ValueError(f"No extractable text found in document: {path.name}")

    full_text = "\n\n".join(paragraphs)
    return [Document(page_content=full_text, metadata={"source_file": path.name, "page": 1})]
