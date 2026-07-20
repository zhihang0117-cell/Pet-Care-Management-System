"""Load Word (.docx) documents from uploads or data/raw."""

from pathlib import Path

from docx import Document as DocxDocument
from langchain_core.documents import Document

from app.config import DATA_RAW

SUPPORTED_EXTENSIONS = {".docx"}


def is_supported_document(filename: str) -> bool:
    return Path(filename).suffix.lower() in SUPPORTED_EXTENSIONS


def load_docx(docx_path: Path) -> list[Document]:
    """Extract text from a Word document as LangChain Documents."""
    doc = DocxDocument(str(docx_path))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    # Include table cell text (policies often use tables in Word)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    paragraphs.append(text)

    if not paragraphs:
        return []

    full_text = "\n\n".join(paragraphs)
    return [
        Document(
            page_content=full_text,
            metadata={
                "source_file": docx_path.name,
                "page": 1,
            },
        )
    ]


def load_document(path: Path) -> list[Document]:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return load_docx(path)
    raise ValueError(f"Unsupported document type: {path.name}. Upload .docx Word files only.")


def load_documents_from_dir(raw_dir: Path | None = None) -> list[Document]:
    raw_dir = raw_dir or DATA_RAW
    all_docs: list[Document] = []
    for doc_path in sorted(raw_dir.glob("*.docx")):
        all_docs.extend(load_docx(doc_path))
    return all_docs
